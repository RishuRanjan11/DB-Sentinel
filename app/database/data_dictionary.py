from __future__ import annotations

import json
import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from app.database.adapters.base import DatabaseAdapter


MAX_DICTIONARY_BYTES = 2 * 1024 * 1024
MAX_DESCRIPTION_LENGTH = 2_000
MAX_TABULAR_ROWS = 10_000

_TABULAR_HEADERS = {
    "schema": {"schema", "schema_name"},
    "table": {"table", "table_name"},
    "column": {
        "column",
        "column_name",
        "field",
        "field_name",
        "variable",
        "variable_name",
    },
    "description": {
        "description",
        "meaning",
        "definition",
        "business_definition",
    },
}


@dataclass(frozen=True)
class DatabaseDictionary:
    description: str
    schema_descriptions: dict[str, str]
    table_descriptions: dict[tuple[str, str], str]
    column_descriptions: dict[tuple[str, str, str], str]

    @classmethod
    def load(
        cls,
        path: str | Path,
        adapter: DatabaseAdapter,
    ) -> DatabaseDictionary:
        dictionary_path = Path(path).expanduser()
        if dictionary_path.stat().st_size > MAX_DICTIONARY_BYTES:
            raise ValueError("Data dictionary files must be 2 MiB or smaller.")

        suffix = dictionary_path.suffix.lower()
        if suffix == ".json":
            try:
                payload = json.loads(
                    dictionary_path.read_text(encoding="utf-8-sig")
                )
            except (OSError, UnicodeError, json.JSONDecodeError) as error:
                raise ValueError(
                    f"Could not read data dictionary JSON: {error}"
                ) from error
        elif suffix == ".csv":
            payload = cls._read_csv(dictionary_path)
        elif suffix in {".xlsx", ".xlsm"}:
            payload = cls._read_excel(dictionary_path)
        else:
            raise ValueError(
                "Data dictionary format must be JSON, CSV, XLSX, or XLSM."
            )

        if not isinstance(payload, dict):
            raise ValueError("Data dictionary root must be a JSON object.")
        if payload.get("version") != 1:
            raise ValueError("Data dictionary must set version to 1.")

        description = cls._description(payload.get("description", ""), "description")
        schemas = adapter.get_schemas()
        actual_schema_names = {
            schema.lower(): schema
            for schema in schemas
        }
        actual_tables = cls._database_tables(adapter, schemas)
        actual_schemas = set(actual_schema_names)
        schema_descriptions: dict[str, str] = {}
        table_descriptions: dict[tuple[str, str], str] = {}
        column_descriptions: dict[tuple[str, str, str], str] = {}
        unknown: list[str] = []
        unscoped_columns = payload.get("unscoped_columns", [])
        if not isinstance(unscoped_columns, list):
            raise ValueError("'unscoped_columns' must be an array.")
        schemas = payload.get("schemas", [])
        if not isinstance(schemas, list):
            raise ValueError("'schemas' must be an array.")

        expanded_schemas = []
        for schema_record in schemas:
            if not isinstance(schema_record, dict):
                raise ValueError("Each schema entry must be a JSON object.")
            if schema_record.get("name"):
                expanded_schemas.append(schema_record)
                continue
            tables = schema_record.get("tables", [])
            if not isinstance(tables, list):
                raise ValueError("Tables for a schema must be an array.")
            for table_record in tables:
                if not isinstance(table_record, dict):
                    raise ValueError("Each table entry must be a JSON object.")
                table_name = cls._required_name(
                    table_record,
                    "name",
                    "table",
                ).lower()
                matches = [
                    key
                    for key in actual_tables
                    if key[1] == table_name
                ]
                if len(matches) != 1:
                    unknown.append(
                        f"{table_name} (schema required because it is "
                        f"{'ambiguous' if matches else 'unknown'})"
                    )
                    continue
                schema_key = matches[0][0]
                expanded_schemas.append({
                    "name": actual_schema_names[schema_key],
                    "tables": [table_record],
                })

        for schema_record in expanded_schemas:
            schema_name = cls._required_name(schema_record, "name", "schema")
            schema_key = schema_name.lower()
            if schema_key not in actual_schemas:
                unknown.append(schema_name)
                continue
            schema_description = cls._description(
                schema_record.get("description", ""),
                f"{schema_name} description",
            )
            if schema_description:
                schema_descriptions[schema_key] = schema_description
            tables = schema_record.get("tables", [])
            if not isinstance(tables, list):
                raise ValueError(f"Tables for schema '{schema_name}' must be a JSON array.")

            for table_record in tables:
                if not isinstance(table_record, dict):
                    raise ValueError("Each table entry must be a JSON object.")
                table_name = cls._required_name(table_record, "name", "table")
                table_key = (schema_key, table_name.lower())
                metadata = actual_tables.get(table_key)
                if metadata is None:
                    unknown.append(f"{schema_name}.{table_name}")
                    continue

                table_description = cls._description(
                    table_record.get("description", ""),
                    f"{schema_name}.{table_name} description",
                )
                if table_description:
                    table_descriptions[table_key] = table_description

                columns = table_record.get("columns", [])
                if not isinstance(columns, list):
                    raise ValueError(
                        f"Columns for '{schema_name}.{table_name}' must be a JSON array."
                    )
                actual_columns = {
                    str(column.get("name", column.get("column", ""))).lower()
                    if isinstance(column, dict)
                    else str(column).lower()
                    for column in metadata.get("columns", [])
                }
                for column_record in columns:
                    if not isinstance(column_record, dict):
                        raise ValueError(
                            f"Column entries for '{schema_name}.{table_name}' must be objects."
                        )
                    column_name = cls._required_name(
                        column_record,
                        "name",
                        f"column in {schema_name}.{table_name}",
                    )
                    if column_name.lower() not in actual_columns:
                        unknown.append(
                            f"{schema_name}.{table_name}.{column_name}"
                        )
                        continue
                    column_description = cls._description(
                        column_record.get("description", ""),
                        f"{schema_name}.{table_name}.{column_name} description",
                    )
                    if column_description:
                        column_descriptions[
                            (schema_key, table_name.lower(), column_name.lower())
                        ] = column_description

        for column_record in unscoped_columns:
            if not isinstance(column_record, dict):
                raise ValueError("Each unscoped column entry must be an object.")
            column_name = cls._required_name(
                column_record,
                "name",
                "column",
            )
            matches = [
                (schema_key, table_key, metadata)
                for (schema_key, table_key), metadata in actual_tables.items()
                if any(
                    (
                        str(column.get("name", column.get("column", "")))
                        if isinstance(column, dict)
                        else str(column)
                    ).lower() == column_name.lower()
                    for column in metadata.get("columns", [])
                )
            ]
            if len(matches) != 1:
                unknown.append(
                    f"{column_name} (column is "
                    f"{'ambiguous' if matches else 'unknown'} without a "
                    "schema/table)"
                )
                continue
            schema_key, table_key, _ = matches[0]
            column_description = cls._description(
                column_record.get("description", ""),
                f"{schema_key}.{table_key}.{column_name} description",
            )
            if column_description:
                column_descriptions[
                    (schema_key, table_key, column_name.lower())
                ] = column_description

        if unknown:
            preview = ", ".join(unknown[:20])
            suffix = " ..." if len(unknown) > 20 else ""
            raise ValueError(
                "Data dictionary references unknown database objects: "
                f"{preview}{suffix}"
            )
        if (
            not description
            and not schema_descriptions
            and not table_descriptions
            and not column_descriptions
        ):
            raise ValueError("Data dictionary contains no descriptions.")

        return cls(
            description=description,
            schema_descriptions=schema_descriptions,
            table_descriptions=table_descriptions,
            column_descriptions=column_descriptions,
        )

    @classmethod
    def _read_csv(cls, path: Path) -> dict[str, Any]:
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.reader(handle)
                headers = next(reader, None)
                if headers is None:
                    raise ValueError("Data dictionary CSV is empty.")
                normalized_headers = [
                    cls._normalize_header(header)
                    for header in headers
                ]
                if len(normalized_headers) != len(set(normalized_headers)):
                    raise ValueError(
                        "Data dictionary CSV has duplicate column headers."
                    )
                rows = []
                for row_number, values in enumerate(reader, start=2):
                    if len(values) != len(headers):
                        raise ValueError(
                            f"Data dictionary CSV row {row_number} has "
                            "a different number of cells than its header."
                        )
                    rows.append(dict(zip(headers, values)))
        except (OSError, UnicodeError, csv.Error) as error:
            raise ValueError(f"Could not read data dictionary CSV: {error}") from error
        return cls._tabular_payload(rows)

    @classmethod
    def _read_excel(cls, path: Path) -> dict[str, Any]:
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError as error:
            raise ValueError(
                "Excel dictionary support requires openpyxl. "
                "Install the project dependencies and try again."
            ) from error

        try:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )
        except (OSError, ValueError, InvalidFileException) as error:
            raise ValueError(
                f"Could not read data dictionary Excel file: {error}"
            ) from error

        try:
            worksheet = workbook.active
            rows_iter = worksheet.iter_rows(values_only=True)
            headers = next(rows_iter, None)
            if headers is None:
                raise ValueError("Data dictionary Excel sheet is empty.")
            normalized_headers = [
                cls._normalize_header(str(header or ""))
                for header in headers
            ]
            if len(normalized_headers) != len(set(normalized_headers)):
                raise ValueError(
                    "Data dictionary Excel sheet has duplicate column headers."
                )
            rows = [
                dict(zip(headers, values))
                for values in rows_iter
            ]
            return cls._tabular_payload(rows)
        finally:
            workbook.close()

    @classmethod
    def _tabular_payload(cls, rows: list[dict[str, Any]]) -> dict[str, Any]:
        if len(rows) > MAX_TABULAR_ROWS:
            raise ValueError(
                f"Tabular data dictionaries may contain at most "
                f"{MAX_TABULAR_ROWS} rows."
            )
        if not rows:
            raise ValueError("Data dictionary file contains no data rows.")

        normalized_rows = []
        for row in rows:
            normalized_rows.append({
                cls._normalize_header(str(header or "")): value
                for header, value in row.items()
                if header is not None
            })

        available_headers = {
            header
            for row in normalized_rows
            for header in row
        }
        selected_headers = {}
        for field, aliases in _TABULAR_HEADERS.items():
            aliases = set(aliases)
            if field == "description":
                aliases.add("definition")
            matches = sorted(available_headers & aliases)
            if len(matches) > 1:
                raise ValueError(
                    f"Data dictionary has multiple columns for '{field}'."
                )
            selected_headers[field] = matches[0] if matches else None

        if not selected_headers["table"] and not selected_headers["schema"]:
            if not selected_headers["column"]:
                raise ValueError(
                    "CSV/Excel data dictionaries without a schema or table "
                    "must identify columns."
                )
        if not selected_headers["description"]:
            raise ValueError(
                "CSV/Excel data dictionaries must include a description column."
            )

        schemas_by_name: dict[str, dict[str, Any]] = {}
        unscoped_columns = []
        database_descriptions = []
        for row_number, row in enumerate(normalized_rows, start=2):
            values = {
                field: row.get(header)
                for field, header in selected_headers.items()
                if header
            }
            schema_name = cls._tabular_text(values.get("schema"))
            table_name = cls._tabular_text(values.get("table"))
            column_name = cls._tabular_text(values.get("column"))
            description = cls._tabular_text(values.get("description"))
            dataset = cls._tabular_text(row.get("dataset"))
            code_list = cls._tabular_text(row.get("code_list"))
            if not any((schema_name, table_name, column_name, description)):
                continue
            if not description:
                raise ValueError(
                    f"Data dictionary row {row_number} has no description."
                )
            description_parts = []
            if dataset:
                description_parts.append(f"Dataset: {dataset}.")
            description_parts.append(description)
            if code_list:
                description_parts[-1] = description_parts[-1].rstrip(".")
                description_parts[-1] += "."
                description_parts.append(f"Code list: {code_list}")
            description = " ".join(description_parts)

            if not schema_name and not table_name:
                if column_name:
                    unscoped_columns.append({
                        "name": column_name,
                        "description": description,
                    })
                    continue
                database_descriptions.append(description)
                continue

            schema_key = schema_name.lower()
            schema_record = schemas_by_name.setdefault(
                schema_key,
                {"name": schema_name, "tables": []},
            )
            if not table_name:
                if column_name:
                    raise ValueError(
                        f"Data dictionary row {row_number} names a column but no table."
                    )
                schema_record["description"] = description
                continue

            table_record = next(
                (
                    item
                    for item in schema_record["tables"]
                    if item["name"].lower() == table_name.lower()
                ),
                None,
            )
            if table_record is None:
                table_record = {"name": table_name, "columns": []}
                schema_record["tables"].append(table_record)
            if not column_name:
                table_record["description"] = description
            else:
                table_record["columns"].append({
                    "name": column_name,
                    "description": description,
                })

        if len(database_descriptions) > 1:
            raise ValueError(
                "Tabular data dictionaries may contain at most one database-level description."
            )
        return {
            "version": 1,
            "description": (
                database_descriptions[0] if database_descriptions else ""
            ),
            "schemas": list(schemas_by_name.values()),
            "unscoped_columns": unscoped_columns,
        }

    @staticmethod
    def _normalize_header(value: str) -> str:
        return re.sub(
            r"[^a-z0-9]+",
            "_",
            value.strip().lower(),
        ).strip("_")

    @staticmethod
    def _tabular_text(value: Any) -> str:
        if value is None:
            return ""
        if not isinstance(value, str):
            raise ValueError(
                "Schema, table, column, and description cells must contain text."
            )
        return value.strip()

    @staticmethod
    def _database_tables(
        adapter: DatabaseAdapter,
        schemas: list[str],
    ) -> dict[tuple[str, str], dict[str, Any]]:
        result = {}
        for schema in schemas:
            for table in adapter.get_tables(schema=schema):
                details = adapter.get_schema(schema=schema, table=table)
                result[(schema.lower(), table.lower())] = details
        return result

    @staticmethod
    def _required_name(record: dict, field: str, label: str) -> str:
        value = record.get(field)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"Every {label} must have a non-empty '{field}'.")
        return value.strip()

    @staticmethod
    def _description(value: Any, label: str) -> str:
        if value is None:
            return ""
        if not isinstance(value, str):
            raise ValueError(f"{label} must be a string.")
        description = value.strip()
        if len(description) > MAX_DESCRIPTION_LENGTH:
            raise ValueError(
                f"{label} must be at most {MAX_DESCRIPTION_LENGTH} characters."
            )
        return description

    def enrich_schema_context(
        self,
        schema_context: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        enriched = []
        for table in schema_context:
            item = dict(table)
            schema = str(item.get("schema", "")).lower()
            table_name = str(item.get("table", "")).lower()
            schema_description = self.schema_descriptions.get(schema)
            if schema_description:
                item["dictionary_schema_description"] = schema_description
            table_description = self.table_descriptions.get(
                (schema, table_name)
            )
            if table_description:
                item["dictionary_description"] = table_description

            columns = []
            for column in item.get("columns", []):
                if not isinstance(column, dict):
                    columns.append(column)
                    continue
                column_item = dict(column)
                column_name = str(
                    column_item.get("name", column_item.get("column", ""))
                ).lower()
                column_description = self.column_descriptions.get(
                    (schema, table_name, column_name)
                )
                if column_description:
                    column_item["dictionary_description"] = column_description
                columns.append(column_item)
            item["columns"] = columns
            enriched.append(item)
        return enriched

    def description_for_prompt(self) -> str:
        table_keys = set(self.table_descriptions)
        table_keys.update(
            (schema, table)
            for schema, table, _ in self.column_descriptions
        )
        records = {
            "database_description": self.description,
            "schemas": self.schema_descriptions,
            "tables": [
                {
                    "schema": schema,
                    "table": table,
                    "description": self.table_descriptions.get(
                        (schema, table), ""
                    ),
                    "columns": {
                        column: column_description
                        for (column_schema, column_table, column), column_description
                        in self.column_descriptions.items()
                        if column_schema == schema and column_table == table
                    },
                }
                for schema, table in sorted(table_keys)
            ],
        }
        return json.dumps(records, ensure_ascii=True, sort_keys=True)

    @property
    def documented_table_count(self) -> int:
        return len(self.table_descriptions)

    @property
    def documented_column_count(self) -> int:
        return len(self.column_descriptions)
