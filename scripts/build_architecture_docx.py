from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ARCHITECTURE_GUIDE.md"
OUTPUT = ROOT / "docs" / "DB-Sentinel-Architecture-Guide.docx"


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(38, 83, 120)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document, lines: list[str]) -> None:
    rows = [
        [value.strip() for value in line.strip().strip("|").split("|")]
        for line in lines
        if line.strip()
    ]
    if len(rows) >= 2 and all(
        re.fullmatch(r":?-{3,}:?", cell.replace(" ", ""))
        for cell in rows[1]
    ):
        rows.pop(1)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[column] if column < len(values) else ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, value)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def build_document() -> None:
    content = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for style_name, size, color in (
        ("Title", 28, "17365D"),
        ("Heading 1", 19, "17365D"),
        ("Heading 2", 14, "1F4E79"),
        ("Heading 3", 11, "365F91"),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("DB-Sentinel | Architecture and Developer Guide").italic = True
    add_page_number(section.footer.paragraphs[0])

    in_code = False
    code_lines: list[str] = []
    index = 0
    first_title = True
    while index < len(content):
        line = content[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(45, 45, 45)
                in_code = False
                code_lines = []
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while index < len(content) and content[index].strip().startswith("|"):
                table_lines.append(content[index])
                index += 1
            add_table(document, table_lines)
            continue

        if not stripped or stripped == "---":
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1 and first_title:
                paragraph = document.add_paragraph(style="Title")
                add_inline(paragraph, title)
                first_title = False
            else:
                document.add_heading(title, level=level)
            index += 1
            continue

        bullet = re.match(r"^\s*-\s+(.*)$", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    if in_code and code_lines:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(8)

    document.core_properties.title = "DB-Sentinel Architecture and Developer Guide"
    document.core_properties.subject = (
        "System overview, source-code guide, data flow, safety and troubleshooting"
    )
    document.core_properties.author = "DB-Sentinel Project"
    document.save(OUTPUT)
    print(f"Generated {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build_document()
